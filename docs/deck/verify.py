"""Re-open the built deck and prove the rules the panel would check.

    tmp/deck/.venv/Scripts/python.exe -m docs.deck.verify [PATH] [--draft]

Exit 1 on any failure. ``--draft`` tolerates [CAPTURE] placeholders and [VERIFY] markers
so an early build can be inspected; the final deck must pass without it.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

from . import theme, xml_post
from .build import OUT_DEFAULT, TOTAL

POOLED_P = '3.222e-05'
PRESENT_P = '0.1257'


def _runs(shape):
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                yield r
    if getattr(shape, 'has_table', False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        yield r


def _text(shape) -> str:
    return ' '.join(r.text for r in _runs(shape))


def verify(path: Path, draft: bool = False) -> list[tuple[str, bool, str]]:
    prs = Presentation(str(path))
    results: list[tuple[str, bool, str]] = []
    slides = list(prs.slides)

    results.append(('slide count', len(slides) == TOTAL, f'{len(slides)} slides'))

    small, missing, placeholders, verify_marks, empty_notes, dup_names, blobs_bad = [], [], [], [], [], [], []
    orb_slides, pooled_alone = [], []
    for idx, slide in enumerate(slides, start=1):
        names = [s.name for s in slide.shapes]
        if len(names) != len(set(names)):
            dup_names.append(f'{idx}: ' + ', '.join(sorted({n for n in names if names.count(n) > 1})))
        if '!!orb' in names:
            orb_slides.append(idx)
        joined = []
        for shape in slide.shapes:
            for r in _runs(shape):
                if not r.text.strip():
                    continue
                if r.font.size is None:
                    missing.append(f'{idx}: {r.text[:40]!r}')
                elif r.font.size < Pt(theme.MIN_PT):
                    small.append(f'{idx}: {r.font.size.pt:.0f}pt {r.text[:40]!r}')
            t = _text(shape)
            joined.append(t)
            if '[CAPTURE]' in t:
                placeholders.append(str(idx))
            if '[VERIFY]' in t:
                verify_marks.append(str(idx))
            if shape.shape_type == 13:  # picture
                rid = shape._element.blipFill.blip.get(qn('r:embed'))
                try:
                    blob = slide.part.related_part(rid).blob
                    if not blob:
                        blobs_bad.append(f'{idx}: empty blob')
                except KeyError:
                    blobs_bad.append(f'{idx}: unresolved {rid}')
        all_text = ' '.join(joined)
        if POOLED_P in all_text and PRESENT_P not in all_text:
            pooled_alone.append(str(idx))
        if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
            empty_notes.append(str(idx))

    results.append((f'every run ≥ {theme.MIN_PT} pt', not small, '; '.join(small[:6]) or 'ok'))
    results.append(('every run has an explicit size', not missing, '; '.join(missing[:6]) or 'ok'))
    results.append(('pooled p never without present p', not pooled_alone, 'slides ' + ', '.join(pooled_alone) if pooled_alone else 'ok'))
    results.append(('pictures resolve', not blobs_bad, '; '.join(blobs_bad[:4]) or 'ok'))
    results.append(('notes on every slide', not empty_notes, 'missing on ' + ', '.join(empty_notes) if empty_notes else 'ok'))
    results.append(('shape names unique per slide', not dup_names, '; '.join(dup_names[:4]) or 'ok'))
    results.append(('!!orb on slides 1, 14, 15', all(n in orb_slides for n in (1, 14, 15)), f'found on {orb_slides}'))
    results.append(('no [CAPTURE] placeholders', draft or not placeholders, 'slides ' + ', '.join(sorted(set(placeholders))) if placeholders else 'ok'))
    results.append(('no [VERIFY] markers', draft or not verify_marks, 'slides ' + ', '.join(sorted(set(verify_marks))) if verify_marks else 'ok'))

    morph_bad = []
    for idx, slide in enumerate(slides, start=1):
        present, dur, opt, fallback = xml_post.has_morph(slide)
        planned = xml_post.MORPH_PLAN.get(idx)
        if planned is None:
            if present:
                morph_bad.append(f'{idx}: unexpected transition')
            continue
        if not present or not fallback or dur != planned[0] or opt != planned[1] or not 500 <= (dur or 0) <= 1500:
            morph_bad.append(f'{idx}: got {present},{dur},{opt},fallback={fallback} want {planned}')
    results.append(('morph plan applied with fallback', not morph_bad, '; '.join(morph_bad[:4]) or 'ok'))

    size = path.stat().st_size
    results.append(('file under 150 MB', size < 150 * 2 ** 20, f'{size / 2 ** 20:.1f} MB'))
    return results


def main(argv=None) -> int:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')   # Windows consoles default to cp1252
    ap = argparse.ArgumentParser()
    ap.add_argument('path', nargs='?', default=str(OUT_DEFAULT))
    ap.add_argument('--draft', action='store_true')
    args = ap.parse_args(argv)
    results = verify(Path(args.path), draft=args.draft)
    width = max(len(name) for name, _, _ in results)
    ok_all = True
    for name, ok, detail in results:
        ok_all &= ok
        print(f'{"PASS" if ok else "FAIL"}  {name:<{width}}  {detail}')
    print('VERIFIED' if ok_all else 'FAILED')
    return 0 if ok_all else 1


if __name__ == '__main__':
    sys.exit(main())
