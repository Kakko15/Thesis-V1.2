"""Thin layer over python-pptx for the deck's visual vocabulary.

Rules baked in here so slide code cannot break them: every run gets an explicit size, face
and colour (the verifier treats a missing size as a failure); text boxes are sized from
measured metrics unless a height is forced, and an over-full forced box raises instead of
overlapping; alpha and shadows are raw DrawingML appended *after* python-pptx styling so
element order stays valid; no grouped shapes (Canva mis-places them on import); persistent
objects are named exactly so Morph matches them across slides. Theme values are read at call
time, never as default arguments, so ``DECK_THEME`` governs everything.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import measure, theme

A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
EMU_IN = 914400


class Overflow(Exception):
    """Text needs more height than the box it was given."""


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def inches(v) -> float:
    return float(v) / EMU_IN


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = theme.SLIDE_W
    prs.slide_height = theme.SLIDE_H
    return prs


def blank_slide(prs: Presentation, ground: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(ground or theme.GROUND)
    return slide


# --- raw DrawingML helpers -----------------------------------------------------------------

def _alpha_on(srgb, pct: float) -> None:
    for old in srgb.findall(qn('a:alpha')):
        srgb.remove(old)
    srgb.append(parse_xml(f'<a:alpha xmlns:a="{A_NS}" val="{int(round(pct * 1000))}"/>'))


def set_fill(shape, hex6: str, alpha_pct: float = 100) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(hex6)
    if alpha_pct < 100:
        srgb = shape._element.spPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        _alpha_on(srgb, alpha_pct)


def set_line(shape, hex6: str, width_pt: float, alpha_pct: float = 100) -> None:
    shape.line.color.rgb = rgb(hex6)
    shape.line.width = Pt(width_pt)
    if alpha_pct < 100:
        srgb = shape._element.spPr.find(qn('a:ln')).find(qn('a:solidFill')).find(qn('a:srgbClr'))
        _alpha_on(srgb, alpha_pct)


def no_line(shape) -> None:
    shape.line.fill.background()


def add_shadow(shape, blur_pt: float = 30, dist_pt: float = 9, alpha_pct: float | None = None) -> None:
    alpha_pct = theme.SHADOW_ALPHA if alpha_pct is None else alpha_pct
    sppr = shape._element.spPr
    for old in sppr.findall(qn('a:effectLst')):
        sppr.remove(old)
    sppr.append(parse_xml(
        f'<a:effectLst xmlns:a="{A_NS}"><a:outerShdw blurRad="{int(blur_pt * 12700)}" '
        f'dist="{int(dist_pt * 12700)}" dir="5400000" algn="t" rotWithShape="0">'
        f'<a:srgbClr val="000000"><a:alpha val="{int(alpha_pct * 1000)}"/></a:srgbClr>'
        f'</a:outerShdw></a:effectLst>'))


def set_picture_opacity(pic, pct: float) -> None:
    blip = pic._element.blipFill.blip
    for old in blip.findall(qn('a:alphaModFix')):
        blip.remove(old)
    blip.append(parse_xml(f'<a:alphaModFix xmlns:a="{A_NS}" amt="{int(pct * 1000)}"/>'))


def set_name(shape, name: str) -> None:
    shape.name = name


# --- primitives --------------------------------------------------------------------------

def glass_card(slide, x, y, w, h, *, name: str | None = None, fill: str | None = None, alpha: float | None = None,
               radius: float = 0.06, border: str | None = None, border_alpha: float | None = None, shadow: bool = True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = radius
    set_fill(shape, fill or theme.PANEL, theme.PANEL_ALPHA if alpha is None else alpha)
    set_line(shape, border or theme.GLASS_BORDER, 0.75, theme.GLASS_BORDER_ALPHA if border_alpha is None else border_alpha)
    if shadow:
        add_shadow(shape)
    shape.text_frame.text = ''
    if name:
        set_name(shape, name)
    return shape


def rect(slide, x, y, w, h, *, fill: str, alpha: float = 100, name: str | None = None, radius: float | None = None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, x, y, w, h)
    if radius is not None:
        shape.adjustments[0] = radius
    set_fill(shape, fill, alpha)
    no_line(shape)
    if name:
        set_name(shape, name)
    return shape


def circle(slide, cx, cy, d, *, fill: str | None, alpha: float = 100, line: str | None = None,
           line_pt: float = 1.5, name: str | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(cx - d / 2)), Emu(int(cy - d / 2)), d, d)
    if fill:
        set_fill(shape, fill, alpha)
    else:
        shape.fill.background()
    if line:
        set_line(shape, line, line_pt)
    else:
        no_line(shape)
    if name:
        set_name(shape, name)
    return shape


def _norm_paragraphs(content, size, font, color, bold):
    """Expand content into paragraphs of fully specified runs."""
    paragraphs = content if isinstance(content, list) else [content]
    out = []
    for para in paragraphs:
        runs = para if isinstance(para, list) else [(para, {})]
        out.append([(t, {
            'size': o.get('size', size), 'font': o.get('font', font),
            'color': o.get('color', color), 'bold': o.get('bold', bold), 'italic': o.get('italic', False),
        }) for t, o in runs])
    return out


def text(slide, x, y, w, h=None, content='', *, size: str | int = 'body', font: str | None = None,
         color: str | None = None, bold: bool = False, align: str = 'left', anchor: str = 'top',
         name: str | None = None, line_spacing: float | None = None, space_after: float | None = None,
         allow_overflow: bool = False):
    """Text box sized from measured metrics.

    ``h=None`` sizes the box to its content. A given ``h`` is checked against the measurement and
    raises :class:`Overflow` when the content would not fit, unless ``allow_overflow``.
    ``content``: str, or list of paragraphs; a paragraph is str or list of ``(text, opts)`` runs
    with opts ``size``, ``font``, ``color``, ``bold``, ``italic``.
    """
    font = font or theme.FONT_BODY
    color = color or theme.TEXT
    size_pt = theme.PT[size] if isinstance(size, str) else max(int(size), theme.MIN_PT)
    paragraphs = _norm_paragraphs(content, size_pt, font, color, bold)
    need = measure.block_height_in(paragraphs, inches(w), space_after_pt=space_after or 0.0, line_spacing=line_spacing)
    if h is None:
        h = Inches(need)
    elif need > inches(h) + 1e-6 and not allow_overflow:
        sample = ' '.join(t for para in paragraphs for t, _ in para)[:60]
        raise Overflow(f'{sample!r} needs {need:.2f} in, box is {inches(h):.2f} in x {inches(w):.2f} in wide')
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}[anchor]
    for i, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
        if line_spacing:
            p.line_spacing = line_spacing
        if space_after is not None:
            p.space_after = Pt(space_after)
        for run_text, o in runs:
            r = p.add_run()
            r.text = run_text
            f = r.font
            s = o['size']
            f.size = Pt(max(theme.PT[s] if isinstance(s, str) else int(s), theme.MIN_PT))
            f.name = o['font']
            f.bold = o['bold']
            f.italic = o['italic']
            f.color.rgb = rgb(o['color'])
    if name:
        set_name(box, name)
    return box


def bottom(shape) -> float:
    """Bottom edge of a shape in inches."""
    return inches(shape.top + shape.height)


def picture(slide, path: Path, x, y, *, w=None, h=None, name: str | None = None, rounded: bool = True,
            border: str | None = 'frame', shadow: bool = True, opacity: float | None = None,
            crop: tuple[float, float, float, float] | None = None):
    pic = slide.shapes.add_picture(str(path), x, y, width=w, height=h)
    if crop:
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    if rounded:
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        av = pic._element.spPr.find(qn('a:prstGeom')).find(qn('a:avLst'))
        for old in av.findall(qn('a:gd')):
            av.remove(old)
        av.append(parse_xml(f'<a:gd xmlns:a="{A_NS}" name="adj" fmla="val 3500"/>'))
    if border:
        set_line(pic, theme.FRAME_BORDER if border == 'frame' else border, 1.25)
    if shadow:
        add_shadow(pic)
    if opacity is not None:
        set_picture_opacity(pic, opacity)
    if name:
        set_name(pic, name)
    return pic


def flat_picture(slide, path: Path, x, y, *, w=None, h=None, name: str | None = None, opacity: float | None = None):
    """A transparent render (orb, glow): no frame, no shadow, no rounding."""
    pic = slide.shapes.add_picture(str(path), x, y, width=w, height=h)
    if opacity is not None:
        set_picture_opacity(pic, opacity)
    if name:
        set_name(pic, name)
    return pic


def chip(slide, x, y, label: str, *, color: str | None = None, filled: bool = False, size: int = 14,
         font: str | None = None, h=Inches(0.36), pad=Inches(0.16), name: str | None = None, ink: str | None = None):
    """Rounded pill with one line of text, width measured from the font."""
    color = color or theme.GREEN
    font = font or theme.FONT_BODY
    size = max(size, theme.MIN_PT)
    w = Inches(measure.width_in(label, font, size, True) + 0.06) + pad * 2
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.5
    if filled:
        set_fill(shape, color, 100)
        set_line(shape, color, 0.75)
        fg = ink or theme.INK_ON_ACCENT
    else:
        set_fill(shape, color, 10 if theme.MODE == 'light' else 12)
        set_line(shape, color, 0.75, 55)
        fg = ink or color
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = pad
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.name = font
    r.font.bold = True
    r.font.color.rgb = rgb(fg)
    if name:
        set_name(shape, name)
    return shape, w


def block_arc(slide, cx, cy, d, *, start_deg: float, sweep_deg: float, thickness: float, fill: str,
              alpha: float = 100, name: str | None = None):
    """A gauge segment. Angles are DrawingML: degrees clockwise from 3 o'clock."""
    shape = slide.shapes.add_shape(MSO_SHAPE.BLOCK_ARC, Emu(int(cx - d / 2)), Emu(int(cy - d / 2)), d, d)
    end_deg = (start_deg + sweep_deg) % 360
    av = shape._element.spPr.find(qn('a:prstGeom')).find(qn('a:avLst'))
    for old in av.findall(qn('a:gd')):
        av.remove(old)
    for gname, val in (('adj1', int(start_deg * 60000)), ('adj2', int(end_deg * 60000)), ('adj3', int(thickness * 100000))):
        av.append(parse_xml(f'<a:gd xmlns:a="{A_NS}" name="{gname}" fmla="val {val}"/>'))
    set_fill(shape, fill, alpha)
    no_line(shape)
    if name:
        set_name(shape, name)
    return shape


def set_notes(slide, notes: str) -> None:
    slide.notes_slide.notes_text_frame.text = notes


def hline(slide, x, y, w, *, color: str | None = None, alpha: float = 12, pt_w: float = 0.75, name: str | None = None):
    conn = slide.shapes.add_connector(1, x, y, x + w, y)
    conn.line.color.rgb = rgb(color or theme.GLASS_BORDER)
    conn.line.width = Pt(pt_w)
    if alpha < 100:
        srgb = conn._element.spPr.find(qn('a:ln')).find(qn('a:solidFill')).find(qn('a:srgbClr'))
        _alpha_on(srgb, alpha)
    if name:
        set_name(conn, name)
    return conn
